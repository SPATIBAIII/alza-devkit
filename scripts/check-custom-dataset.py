"""Real browser acceptance check; run with localhost API/web and Edge installed."""
import json
from io import BytesIO
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


with sync_playwright() as playwright:
    browser = playwright.chromium.launch(channel="msedge")
    page = browser.new_page(viewport={"width": 1440, "height": 1000})
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto("http://localhost:3000")
    page.get_by_role("button", name="Mock Data Factory", exact=True).click()
    page.get_by_role("button", name="Custom Dataset", exact=True).click()
    page.get_by_label("Dataset name", exact=True).fill("Browser Student Enrollment")
    page.get_by_role("button", name="Add field", exact=True).click()
    expect(page.get_by_role("article")).to_have_count(5)
    page.get_by_label("Field 5 type", exact=True).select_option("integer")
    last = page.get_by_role("article").nth(4)
    last.get_by_text("Advanced settings", exact=True).click()
    expect(last.get_by_label("field_1 Minimum", exact=True)).to_be_visible()
    expect(last.get_by_text("Minimum length", exact=True)).to_have_count(0)
    page.get_by_role("button", name="Delete field_1", exact=True).click()
    expect(page.get_by_role("article")).to_have_count(4)
    page.get_by_role("button", name="Duplicate student_name", exact=True).click()
    page.get_by_role("button", name="Move student_name_1 up", exact=True).click()
    expect(page.get_by_label("Field 2 name", exact=True)).to_have_value("student_name_1")
    page.get_by_role("button", name="Delete student_name_1", exact=True).click()
    page.get_by_label("Field 2 name", exact=True).fill("student_id")
    page.get_by_role("button", name="Validate Schema", exact=True).click()
    expect(page.locator("main [role=alert]")).to_contain_text("Field names must be unique")
    page.get_by_label("Field 2 name", exact=True).fill("student_name")
    page.get_by_role("button", name="Validate Schema", exact=True).click()
    expect(page.get_by_text("Schema and generation constraints are valid.", exact=True)).to_be_visible()
    page.get_by_role("button", name="Generate Dataset", exact=True).click()
    expect(page.get_by_text("100 synthetic records · Seed 42", exact=True)).to_be_visible()
    expect(page.locator("table:visible th")).to_have_count(4)
    expect(page.locator("table:visible tbody tr")).to_have_count(25)
    page.get_by_role("button", name="Next", exact=True).click()
    expect(page.get_by_text("STU-00026", exact=True)).to_be_visible()
    page.get_by_role("button", name="Previous", exact=True).click()
    first_name = page.locator("table:visible tbody tr").first.locator("td").nth(1).inner_text()
    for label, extension in [("CSV", "csv"), ("JSON", "json"), ("Excel", "xlsx")]:
        with page.expect_download() as info:
            page.get_by_role("button", name=f"Download {label}", exact=True).click()
        download = info.value
        content = open(download.path(), "rb").read()
        assert download.suggested_filename.endswith(extension)
        if extension == "json":
            assert json.loads(content)["data"][0]["student_name"] == first_name
        elif extension == "xlsx":
            workbook = load_workbook(BytesIO(content))
            assert workbook.active.max_row == 101
            assert workbook.active["B2"].value == first_name
            workbook.close()
    page.get_by_label("Dataset name", exact=True).fill("Browser Student Enrollment changed")
    expect(page.get_by_text("Preview is outdated.", exact=False)).to_be_visible()
    expect(page.get_by_role("button", name="Download Excel", exact=True)).to_be_disabled()
    page.get_by_role("button", name="Save Template", exact=True).click()
    expect(page.get_by_text("Template saved.", exact=True)).to_be_visible()
    saved_id = page.get_by_label("Load template", exact=True).locator('optgroup[label="Saved templates"] option').filter(has_text="Browser Student Enrollment changed").get_attribute("value")
    page.reload()
    page.get_by_role("button", name="Mock Data Factory", exact=True).click()
    page.get_by_role("button", name="Custom Dataset", exact=True).click()
    page.get_by_label("Load template", exact=True).select_option(saved_id)
    expect(page.get_by_label("Dataset name", exact=True)).to_have_value("Browser Student Enrollment changed")
    page.get_by_role("button", name="Generate Dataset", exact=True).click()
    expect(page.get_by_text("100 synthetic records · Seed 42", exact=True)).to_be_visible()
    with page.expect_download() as info:
        page.get_by_role("button", name="Export Template JSON", exact=True).click()
    template_path = info.value.path()
    page.get_by_label("Import template JSON", exact=False).set_input_files(template_path)
    expect(page.get_by_text("Template imported.", exact=True)).to_be_visible()
    page.get_by_role("button", name="Delete Template", exact=True).click()
    expect(page.get_by_text("Template deleted.", exact=True)).to_be_visible()
    page.get_by_label("Load template", exact=True).select_option(saved_id)
    expect(page.get_by_role("button", name="Duplicate Template", exact=True)).to_be_enabled()
    page.get_by_role("button", name="Duplicate Template", exact=True).click()
    expect(page.get_by_text("Saved template duplicated.", exact=True)).to_be_visible()
    page.get_by_role("button", name="Delete Template", exact=True).click()
    expect(page.get_by_text("Template deleted.", exact=True)).to_be_visible()
    page.get_by_label("Load template", exact=True).select_option(saved_id)
    expect(page.get_by_role("button", name="Delete Template", exact=True)).to_be_enabled()
    page.set_viewport_size({"width": 390, "height": 844})
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    page.get_by_role("button", name="Delete Template", exact=True).click()
    expect(page.get_by_text("Template deleted.", exact=True)).to_be_visible()
    page.get_by_role("button", name="Household Dataset", exact=True).click()
    page.get_by_role("button", name="Generate Dataset", exact=True).click()
    expect(page.get_by_text("100 synthetic records", exact=False)).to_be_visible()
    assert not errors, errors
    browser.close()
    print("Custom dataset browser checks passed.")
