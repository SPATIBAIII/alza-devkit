"""Browser integration check. Run with both servers running; requires Playwright and Edge."""

import csv
import json
from io import BytesIO, StringIO

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright


with sync_playwright() as playwright:
    browser = playwright.chromium.launch(channel="msedge")
    page = browser.new_page(viewport={"width": 1440, "height": 1000})
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto("http://localhost:3000")
    expect(page.get_by_role("heading", name="Developer Tools")).to_be_visible()
    page.get_by_role("button", name="Mock Data Factory", exact=True).click()
    expect(page.get_by_text("Choose your settings", exact=False)).to_be_visible()
    page.get_by_label("Record count").fill("1001")
    page.get_by_role("button", name="Generate Dataset", exact=True).click()
    expect(page.get_by_text("Enter a whole number from 1 to 1,000.")).to_be_visible()
    page.get_by_label("Record count").fill("100")
    page.get_by_role("button", name="Generate Dataset", exact=True).click()
    expect(page.get_by_text("100 synthetic records", exact=False)).to_be_visible()
    expect(page.locator("tbody tr")).to_have_count(25)
    first_name = page.locator("tbody tr").first.locator("td").nth(1).inner_text()
    page.get_by_role("button", name="Next", exact=True).click()
    expect(page.get_by_text("HH-00026", exact=True)).to_be_visible()
    page.get_by_role("button", name="Previous", exact=True).click()

    # Unsaved form edits must not affect downloads of the displayed dataset.
    page.get_by_label("Record count").fill("5")
    page.get_by_label("Random seed").fill("7")
    for label, extension in [("CSV", "csv"), ("JSON", "json"), ("Excel", "xlsx")]:
        with page.expect_download() as download_info:
            page.get_by_role("button", name=f"Download {label}", exact=True).click()
        download = download_info.value
        assert download.suggested_filename == f"synthetic-households-100-seed-42.{extension}"
        with open(download.path(), "rb") as file:
            content = file.read()
        if extension == "json":
            dataset = json.loads(content)
            assert dataset["total_records"] == 100 and dataset["seed"] == 42
            assert dataset["data"][0]["household_head"] == first_name
        elif extension == "csv":
            records = list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
            assert len(records) == 100 and records[0]["household_head"] == first_name
        else:
            workbook = load_workbook(BytesIO(content))
            assert workbook.active.max_row == 101
            assert workbook.active["B2"].value == first_name
            workbook.close()

    page.get_by_role("button", name="Next", exact=True).click()
    page.get_by_role("button", name="Generate Dataset", exact=True).click()
    expect(page.get_by_text("5 synthetic records", exact=False)).to_be_visible()
    expect(page.get_by_text("Page 1 of 1")).to_be_visible()
    expect(page.locator("tbody tr")).to_have_count(5)

    page.route("**/api/v1/data-factory/**", lambda route: route.abort())
    page.get_by_role("button", name="Generate Dataset", exact=True).click()
    expect(page.locator("main [role=alert]")).to_contain_text("Could not reach the data service")
    expect(page.locator("tbody tr")).to_have_count(5)
    page.get_by_role("button", name="Download CSV", exact=True).click()
    expect(page.locator("main [role=alert]").last).to_contain_text("Could not reach the data service")
    page.unroute("**/api/v1/data-factory/**")

    page.set_viewport_size({"width": 390, "height": 844})
    page.get_by_role("button", name="Open navigation", exact=True).click()
    page.get_by_role("button", name="Overview", exact=True).click()
    expect(page.get_by_role("heading", name="Developer Tools")).to_be_visible()
    page.get_by_role("button", name="Open navigation", exact=True).click()
    page.get_by_role("button", name="Mock Data Factory", exact=True).click()
    page.get_by_role("button", name="Generate Dataset", exact=True).click()
    expect(page.get_by_text("100 synthetic records", exact=False)).to_be_visible()
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    region = page.get_by_role("region", name="Synthetic household records")
    assert region.evaluate("el => el.scrollWidth > el.clientWidth")
    assert errors == [], errors
    browser.close()
    print("PASS: browser generation, validation, pagination, all downloads, config snapshot, error recovery, overview and mobile layout")
