import json
import re
from io import BytesIO
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.modules.data_factory.exporters import export_records, spreadsheet_safe


def export_dataset(dataset, format):
    payload = dataset.model_dump()
    if format == "json":
        output = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    elif format == "csv":
        stream = BytesIO()
        with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
            files = {}
            for index, key in enumerate(dataset.order, 1):
                slug = re.sub(r"[^a-zA-Z0-9_-]", "_", key)[:70]
                filename = f"{index:02d}-{slug}.csv"
                files[key] = filename
                archive.writestr(ZipInfo(filename), export_records({"data": dataset.tables[key]}, dataset.fields[key], "csv"), compress_type=ZIP_DEFLATED)
            manifest = {k: v for k, v in payload.items() if k != "tables"}
            manifest["files"] = files
            archive.writestr(ZipInfo("manifest.json"), json.dumps(manifest, ensure_ascii=False, indent=2), compress_type=ZIP_DEFLATED)
        output = stream.getvalue()
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for index, key in enumerate(dataset.order, 1):
            # Numeric prefix ensures uniqueness even after truncation/case folding.
            title = f"{index:02d}-" + re.sub(r"[\\/*?:\[\]'\x00-\x1f]", "_", key)
            sheet = workbook.create_sheet(title[:31])

            def safe(value):
                return spreadsheet_safe(ILLEGAL_CHARACTERS_RE.sub("", value).replace("\r\n", "\n").replace("\r", "\n")) if isinstance(value, str) else value

            sheet.append([safe(c) for c in dataset.fields[key]])
            for row in dataset.tables[key]:
                sheet.append([safe(row[c]) for c in dataset.fields[key]])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        stream = BytesIO()
        try:
            workbook.save(stream)
            output = stream.getvalue()
        finally:
            workbook.close()
    if len(output) > 25_000_000:
        raise ValueError("Export exceeds 25 MB output limit")
    return output
