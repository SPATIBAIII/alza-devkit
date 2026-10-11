import csv
import json
from io import BytesIO, StringIO
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.main import app
from app.modules.schema_import.parser import parse_sql
from app.modules.schema_import.models import DatabaseSchema, identity
from app.modules.schema_import.generator import generate
from app.modules.schema_import.store import SchemaStore
from app.modules.schema_import.exporters import export_dataset
from app.modules.data_factory.template_store import TemplateStore
from app.modules.data_factory.custom_schemas import TemplateInput

BASE = "/api/v1/schema-import"
STUDENTS = '''
CREATE TABLE IF NOT EXISTS public.students (
 id SERIAL PRIMARY KEY, first_name VARCHAR(80) NOT NULL,
 email TEXT UNIQUE, birth_date DATE, active BOOLEAN DEFAULT true
);
CREATE TABLE public.enrollments (
 id BIGSERIAL PRIMARY KEY, student_id INTEGER NOT NULL,
 quantity INTEGER CHECK (quantity >= 0), price NUMERIC(8,2),
 FOREIGN KEY (student_id) REFERENCES public.students(id) ON DELETE CASCADE ON UPDATE RESTRICT
);
'''


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_FACTORY_DB", str(tmp_path / "templates.sqlite3"))
    with TestClient(app) as client:
        yield client


def test_basic_parse_mapping_and_grouped_constraints():
    schema = parse_sql(STUDENTS)
    parent, child = schema.tables
    assert parent.key == '"public"."students"'
    assert parent.primary_key == ["id"]
    assert not parent.columns[0].nullable
    assert parent.unique_constraints == [["email"]]
    assert [c.synthetic_generator for c in parent.columns] == ["sequence", "first_name", "email", "birth_date", "boolean"]
    assert child.columns[-1].synthetic_generator == "price"
    assert child.foreign_keys[0].on_delete == "CASCADE"
    assert child.foreign_keys[0].on_update == "RESTRICT"
    assert parent.columns[-1].default == "true"
    assert any("never evaluated" in w for w in schema.warnings)


def test_quoted_names_comments_literals_and_schema_collision():
    schema = parse_sql('''-- comment ; CREATE TABLE evil
      /* nested /* ; */ comment */
      CREATE TABLE "odd.schema"."Table.Name" ("a""b" INTEGER PRIMARY KEY, "email;" TEXT DEFAULT 'x,'';--y');
      CREATE TABLE other."Table.Name" (id INTEGER);
      CREATE TABLE public."Table.Name" (id INTEGER);
    ''')
    assert len({t.key for t in schema.tables}) == 3
    assert schema.tables[0].columns[0].name == 'a"b'
    assert schema.tables[0].columns[1].default == "'x,'';--y'"
    generate(schema)


@pytest.mark.parametrize("sql_type", ["INTEGER", "BIGINT", "SMALLINT", "VARCHAR(20)", "CHAR(5)", "TEXT", "BOOLEAN", "DATE", "TIMESTAMP", "TIMESTAMPTZ", "UUID", "NUMERIC(7,2)", "DECIMAL", "REAL", "DOUBLE PRECISION", "SERIAL", "BIGSERIAL", "TIMESTAMP(3) WITH TIME ZONE", "TIMESTAMP WITHOUT TIME ZONE", "CHARACTER VARYING(20)"])
def test_sql_types(sql_type):
    schema = parse_sql(f"CREATE TABLE t (v {sql_type})")
    schema.tables[0].row_count = 5
    assert len(generate(schema).tables[schema.tables[0].key]) == 5


def test_enum_and_default_casts():
    schema = parse_sql("""CREATE TYPE public.status AS ENUM ('pending','done','it''s fine');
      CREATE TABLE t (id UUID DEFAULT gen_random_uuid(), state public.status NOT NULL DEFAULT 'pending'::public.status,
      created TIMESTAMP DEFAULT now(), n INT DEFAULT nextval('t_seq'::regclass));""")
    col = schema.tables[0].columns[1]
    assert col.enum_values == ["pending", "done", "it's fine"]
    assert col.synthetic_generator == "enum"
    assert all(row["state"] in col.enum_values for row in generate(schema).tables[schema.tables[0].key])


@pytest.mark.parametrize("sql", [
    "", "SELECT 1", "DROP TABLE students", "CREATE VIEW a AS SELECT 1", "CREATE TABLE t (id jsonb)",
    "CREATE TABLE t (id INT", 'CREATE TABLE "missing (id INT)', "/* unterminated", "CREATE TABLE t (v TEXT DEFAULT 'oops)",
    "CREATE TABLE t (a INT,)", "CREATE TABLE t (a INT a INT)", "CREATE TABLE t (a INT, a INT)",
    "CREATE TABLE t (a INT) CREATE TABLE u (b INT)", "CREATE TABLE t (a INT) PARTITION BY RANGE (a)",
    "CREATE TABLE t (a INT GENERATED ALWAYS AS IDENTITY)", "CREATE TABLE t (a INT CHECK (a > 0 AND a < 10))",
    "CREATE TABLE t (a TEXT CHECK (length(a) > 1))", "CREATE TABLE t (a TEXT COLLATE c)",
    "CREATE TABLE t (a INT DEFAULT 1 COLLATE c)", "CREATE TABLE t (a INT DEFAULT nonsense())",
    "CREATE TABLE t (a INT DEFAULT 1 + 2)", "CREATE TABLE t (a INT[])", "CREATE TABLE t (a INT); DELETE FROM t;",
    "CREATE TABLE t (a INT NOT NULL NULL)", "CREATE TABLE t (a INT NULL NOT NULL)",
    "CREATE TABLE t (a INT PRIMARY KEY, b INT PRIMARY KEY)", "CREATE TABLE t (a INT, UNIQUE NULLS NOT DISTINCT(a))",
    "CREATE TABLE t (a INT REFERENCES t(a) DEFERRABLE)", "CREATE TABLE t (a INT, PRIMARY KEY(b))",
    "CREATE TABLE t (a VARCHAR(0))", "CREATE TABLE t (a DECIMAL(2,3))", "CREATE TABLE t (a INT(5))",
    "CREATE TABLE t (a TEXT DEFAULT $$secret$$)", "CREATE TABLE t (a TEXT DEFAULT E'escaped')",
    "CREATE TABLE t (a INT); ALTER TABLE t DROP COLUMN a", "CREATE TYPE s AS ENUM ('a','a')",
])
def test_reject_malformed_and_unsupported(sql):
    with pytest.raises((ValueError, ValidationError)):
        parse_sql(sql)


@pytest.mark.parametrize("sql", [
    "--" + "x" * 262144,
    "--" + "é" * 131072,
    "CREATE TABLE t (v INT DEFAULT " + "(" * 33 + "1" + ")" * 33 + ")",
    "/*" * 33 + "*/" * 33,
    "CREATE TABLE t (" + ",".join(f"c{i} INT" for i in range(51)) + ")",
    ";".join(f"CREATE TABLE t{i} (id INT)" for i in range(21)),
], ids=["byte-limit", "utf8-limit", "parenthesis-depth", "comment-depth", "column-limit", "table-limit"])
def test_parser_limits(sql):
    with pytest.raises(ValueError):
        parse_sql(sql)


def test_acceptance_100_students_300_enrollments_seed_and_fk_integrity():
    schema = parse_sql(STUDENTS)
    schema.tables[1].row_count = 300
    result = generate(schema)
    assert result == generate(schema)
    parent, child = [result.tables[k] for k in result.order]
    assert len(parent) == 100 and len(child) == 300
    ids = {row["id"] for row in parent}
    assert all(row["student_id"] in ids for row in child)
    assert len({row["id"] for row in child}) == 300
    schema.seed = 21
    assert result != generate(schema)


def test_composite_fk_uses_tuples_and_alter_constraints():
    schema = parse_sql('''CREATE TABLE p (a INT, b INT);
      CREATE TABLE c (id SERIAL PRIMARY KEY, a INT, b INT);
      ALTER TABLE p ADD CONSTRAINT p_pk PRIMARY KEY (a,b);
      ALTER TABLE c ADD CONSTRAINT c_fk FOREIGN KEY (a,b) REFERENCES p(a,b);
      ALTER TABLE c ADD CONSTRAINT c_uniq UNIQUE (id,a);
    ''')
    schema.tables[0].columns[0].synthetic_generator = "integer"
    schema.tables[0].columns[0].generator_options.min = 0
    schema.tables[0].columns[0].generator_options.max = 9
    schema.tables[0].columns[1].synthetic_generator = "integer"
    schema.tables[0].columns[1].generator_options.min = 0
    schema.tables[0].columns[1].generator_options.max = 9
    schema.tables[0].row_count = 50
    schema.tables[1].row_count = 300
    result = generate(schema)
    p, c = [result.tables[k] for k in result.order]
    pairs = {(r["a"], r["b"]) for r in p}
    assert len(pairs) == 50
    assert len({r["a"] for r in p}) < 50  # not independently unique
    assert all((r["a"], r["b"]) in pairs for r in c)


@pytest.mark.parametrize("ddl, message", [
    ("CREATE TABLE a (id INT PRIMARY KEY, parent INT REFERENCES a(id))", "Cyclic"),
    ("CREATE TABLE a (id INT PRIMARY KEY, b INT REFERENCES b(id)); CREATE TABLE b (id INT PRIMARY KEY, a INT REFERENCES a(id))", "Cyclic"),
    ("CREATE TABLE a (id INT REFERENCES missing(id))", "missing table"),
    ("CREATE TABLE a (id INT); CREATE TABLE b (id INT REFERENCES a(id))", "declared primary or unique"),
    ("CREATE TABLE a (id INT PRIMARY KEY); CREATE TABLE b (id TEXT REFERENCES a(id))", "incompatible"),
    ("CREATE TABLE a (id INT PRIMARY KEY); CREATE TABLE b (x INT, y INT, FOREIGN KEY(x,y) REFERENCES a(id))", "equal lengths"),
    ("CREATE TABLE a (id INT PRIMARY KEY); CREATE TABLE b (x INT REFERENCES a(id), FOREIGN KEY(x) REFERENCES a(id))", "Overlapping"),
    ("CREATE TABLE a (id BOOLEAN UNIQUE)", "capacity"),
    ("CREATE TABLE a (id INT CHECK (id > 200))", "CHECK"),
    ("CREATE TABLE a (id TEXT CHECK (id > 1))", "numeric column"),
])
def test_relationship_and_constraint_failures(ddl, message):
    with pytest.raises(ValueError, match=message):
        generate(parse_sql(ddl))


def test_selection_counts_and_capacity():
    schema = parse_sql("CREATE TABLE p (id INT PRIMARY KEY); CREATE TABLE c (p INT UNIQUE REFERENCES p);")
    schema.tables[0].row_count = 10
    schema.tables[1].row_count = 10
    data = generate(schema)
    assert len({r["p"] for r in data.tables[schema.tables[1].key]}) == 10
    schema.tables[1].row_count = 11
    with pytest.raises(ValueError, match="capacity"):
        generate(schema)
    schema.tables[0].row_count = 0
    with pytest.raises(ValueError, match="at least one"):
        generate(schema)
    schema.tables[0].selected = False
    with pytest.raises(ValueError, match="Select parent"):
        generate(schema)
    schema.tables[1].selected = False
    with pytest.raises(ValueError, match="Select at least"):
        generate(schema)


def test_zero_rows_and_total_limit():
    schema = parse_sql(";".join(f"CREATE TABLE t{i} (id INT)" for i in range(6)))
    for t in schema.tables:
        t.row_count = 1000
    with pytest.raises(ValueError, match="5,000"):
        generate(schema)
    for t in schema.tables:
        t.row_count = 0
    assert generate(schema).total_records == 0


def test_check_enforcement_and_mapping_overrides():
    schema = parse_sql("CREATE TABLE t (id SERIAL PRIMARY KEY, quantity INT CHECK (quantity >= 90), amount DECIMAL(3,2))")
    schema.tables[0].row_count = 30
    rows = generate(schema).tables[schema.tables[0].key]
    assert all(r["quantity"] >= 90 and 0 <= r["amount"] < 10 for r in rows)
    payload = schema.model_dump()
    payload["tables"][0]["columns"][1]["synthetic_generator"] = "enum"
    payload["tables"][0]["columns"][1]["generator_options"]["values"] = ["not a number"]
    with pytest.raises(ValueError, match="mapping cannot satisfy"):
        generate(payload)


def test_exports_formula_protection_and_stable_order():
    schema = parse_sql('''CREATE TABLE "=unsafe" ("=header" TEXT, id INT PRIMARY KEY);
      CREATE TABLE second (id INT PRIMARY KEY);''')
    schema.tables[0].columns[0].synthetic_generator = "enum"
    schema.tables[0].columns[0].generator_options.values = ["\t=HYPERLINK(\"x\")"]
    data = generate(schema)
    exported = json.loads(export_dataset(data, "json"))
    assert exported["tables"][data.order[0]][0]["=header"].startswith("\t=")
    csv_bytes = export_dataset(data, "csv")
    assert csv_bytes == export_dataset(data, "csv")
    with ZipFile(BytesIO(csv_bytes)) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest["files"]) == 2
        assert manifest["order"] == data.order
        for filename in archive.namelist():
            assert "/" not in filename and "\\" not in filename
        rows = list(csv.reader(StringIO(archive.read(manifest["files"][data.order[0]]).decode("utf-8-sig"))))
        assert rows[0] == ["'=header", "id"]
        assert rows[1][0].startswith("'\t=")
    workbook = load_workbook(BytesIO(export_dataset(data, "xlsx")))
    assert len(workbook.worksheets) == 2
    assert workbook.worksheets[0]["A1"].value == "'=header"
    assert workbook.worksheets[0]["A2"].data_type == "s"
    assert workbook.worksheets[0].max_row == 101
    workbook.close()


def test_storage_restart_roundtrip_and_template_isolation(tmp_path):
    path = tmp_path / "shared.sqlite3"
    original = TemplateStore(path).create(TemplateInput(name="Original", fields=[{"name": "v", "type": "integer"}]))
    store = SchemaStore(path)
    saved = store.create(parse_sql(STUDENTS))
    restarted = SchemaStore(path)
    assert restarted.get(saved.id) == saved
    assert generate(restarted.get(saved.id)) == generate(saved)
    imported = restarted.create(DatabaseSchema.model_validate_json(saved.model_dump_json()))
    assert imported.id != saved.id
    saved.name = "Renamed"
    assert restarted.update(saved.id, saved).name == "Renamed"
    restarted.delete(imported.id)
    assert len(restarted.list()) == 1
    assert TemplateStore(path).get(original.id) == original


def test_api_workflow(client):
    response = client.post(BASE + "/parse", json={"sql": STUDENTS, "name": "Acceptance"})
    assert response.status_code == 200, response.text
    schema = response.json()
    schema["tables"][1]["row_count"] = 300
    validation = client.post(BASE + "/validate", json=schema)
    assert validation.status_code == 200
    data = client.post(BASE + "/generate", json=schema)
    assert data.status_code == 200 and data.json()["total_records"] == 400
    saved = client.post(BASE + "/schemas", json=schema)
    assert saved.status_code == 201
    id = saved.json()["id"]
    assert client.get(BASE + "/schemas").json()[0]["id"] == id
    exported = client.get(f"{BASE}/schemas/{id}/export")
    imported = client.post(BASE + "/schemas/import", json=exported.json())
    assert imported.status_code == 201 and imported.json()["id"] != id
    schema["name"] = "Renamed"
    assert client.put(f"{BASE}/schemas/{id}", json=schema).json()["name"] == "Renamed"
    for format, media in [("json", "application/json"), ("csv", "application/zip"), ("xlsx", "application/vnd.openxmlformats")]:
        response = client.post(BASE + "/export?format=" + format, json=schema)
        assert response.status_code == 200 and response.headers["content-type"].startswith(media)
    assert client.delete(f"{BASE}/schemas/{id}").status_code == 204
    assert client.get(f"{BASE}/schemas/{id}").status_code == 404


@pytest.mark.parametrize("change", ["count", "seed", "version", "credentials", "type", "unknown-column"])
def test_api_rejects_invalid_configs(client, change):
    schema = parse_sql(STUDENTS).model_dump(mode="json")
    if change == "count":
        schema["tables"][0]["row_count"] = 1001
    elif change == "seed":
        schema["seed"] = -1
    elif change == "version":
        schema["version"] = 99
    elif change == "credentials":
        schema["password"] = "never-store-this"
    elif change == "type":
        schema["tables"][0]["columns"][0]["synthetic_generator"] = "eval"
    else:
        schema["tables"][0]["primary_key"] = ["unknown"]
    assert client.post(BASE + "/generate", json=schema).status_code == 422


def test_api_payload_and_parser_errors(client):
    assert client.post(BASE + "/parse", json={"sql": "SELECT secret"}).status_code == 422
    assert client.post(BASE + "/parse", json={"sql": "--" + "é" * 140000}).status_code == 422
    response = client.post(BASE + "/parse", content=b" " * 2_000_001, headers={"content-type": "application/json"})
    assert response.status_code == 413
    assert "Traceback" not in response.text
    assert client.post(BASE + "/parse", content=b'{"sql":').status_code == 422
