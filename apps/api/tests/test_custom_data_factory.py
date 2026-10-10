import csv
import json
import random
from io import BytesIO, StringIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.main import app
from app.modules.data_factory.custom_generator import generate_custom
from app.modules.data_factory.custom_schemas import GenerateRequest, TemplateInput
from app.modules.data_factory.field_registry import REGISTRY
from app.modules.data_factory.template_store import TemplateStore

BASE = "/api/v1/data-factory"


def field(name="value", type="string", **kwargs):
    return {"name": name, "type": type, **kwargs}


def request(fields=None, count=10, seed=42):
    return {"template": {"name": "Test dataset", "fields": fields or [field()]}, "count": count, "seed": seed}


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_FACTORY_DB", str(tmp_path / "templates.sqlite3"))
    with TestClient(app) as client:
        yield client


@pytest.mark.parametrize("kind", sorted(REGISTRY))
def test_every_field_type(kind):
    options = {"values": ["A", "B"]} if kind == "enum" else {}
    payload = request([field(type=kind, options=options)])
    result = generate_custom(payload)
    assert result.total_records == len(result.data) == 10
    assert all(row["value"] is not None for row in result.data)
    assert result == generate_custom(payload)


def test_local_random_and_seeds():
    before = random.getstate()
    assert generate_custom(request()) != generate_custom(request(seed=7))
    assert random.getstate() == before


@pytest.mark.parametrize("configuration", [
    field(type="invalid"), field(name="bad name"), field(options={"min": 10, "max": 0}),
    field(type="integer", options={"min": 0.5}), field(type="enum"),
    field(type="enum", options={"values": ["A", "A"]}),
    field(type="enum", options={"values": ["A"], "weights": [0]}),
    field(type="enum", options={"values": ["A"], "weights": [1, 2]}),
    field(options={"min_date": "2025-01-01", "max_date": "2020-01-01"}),
    field(options={"max_length": 500}), field(null_probability=0.5),
    field(type="boolean", default="false"), field(type="integer", default=1000),
    field(type="enum", options={"values": ["A"]}, default="B"),
    field(type="computed", computed={"operation": "eval", "sources": ["x"]}),
    field(options={"code": "print('hello')"}),
])
def test_invalid_fields(configuration):
    with pytest.raises(ValidationError):
        GenerateRequest.model_validate(request([configuration]))


@pytest.mark.parametrize("count", [0, -1, 1001, True, 1.5])
def test_invalid_count(client, count):
    assert client.post(f"{BASE}/custom/generate", json=request(count=count)).status_code == 422


def test_nulls_and_defaults():
    result = generate_custom(request([field(required=False, null_probability=1), field("fixed", "integer", default=7)]))
    assert all(r == {"value": None, "fixed": 7} for r in result.data)
    values = generate_custom(request([field(required=False, null_probability=0.5)], count=100)).data
    assert any(r["value"] is None for r in values) and any(r["value"] is not None for r in values)


def test_boundaries_and_full_unique_domain():
    data = generate_custom(request([field(type="integer", unique=True, options={"min": 1, "max": 1000})], count=1000)).data
    assert {r["value"] for r in data} == set(range(1, 1001))
    data = generate_custom(request([field(type="decimal", options={"min": 1.23, "max": 1.23, "precision": 2})])).data
    assert all(r["value"] == 1.23 for r in data)


@pytest.mark.parametrize("configuration", [field(type="boolean", unique=True), field(type="integer", unique=True, options={"min": 1, "max": 2}), field(type="string", unique=True, default="fixed"), field(type="enum", unique=True, options={"values": ["A", "B"], "weights": [0, 1]})])
def test_impossible_unique(client, configuration):
    response = client.post(f"{BASE}/custom/validate", json=request([configuration]))
    assert response.status_code == 422
    assert "domain" in response.json()["detail"]


def test_weighted_enum_unique():
    result = generate_custom(request([field(type="enum", unique=True, options={"values": ["A", "B", "C"], "weights": [10, 1, 0]})], count=2))
    assert {r["value"] for r in result.data} == {"A", "B"}


def test_computed_and_field_order():
    fields = [
        field("total", "computed", computed={"operation": "multiply", "sources": ["quantity", "price"]}),
        field("quantity", "quantity", default=3), field("price", "price", default=2.5),
        field("first", "first_name", default="Ana"), field("last", "last_name", default="Sample"),
        field("name", "computed", computed={"operation": "full_name", "sources": ["first", "last"]}),
        field("birth", "birth_date", options={"min_date": "2000-03-01", "max_date": "2000-03-01"}),
        field("age", "computed", computed={"operation": "age", "sources": ["birth"], "reference_date": "2026-02-28"}),
    ]
    result = generate_custom(request(fields))
    assert list(result.data[0]) == [f["name"] for f in fields]
    assert all(r["total"] == 7.5 and r["name"] == "Ana Sample" and r["age"] == 25 for r in result.data)


@pytest.mark.parametrize("fields", [
    [field("x"), field("x")],
    [field("x", "computed", computed={"operation": "multiply", "sources": ["y", "y"]})],
    [field("x", "computed", computed={"operation": "multiply", "sources": ["x", "x"]})],
    [field("x", "computed", computed={"operation": "multiply", "sources": ["y", "y"]}), field("y", "computed", computed={"operation": "multiply", "sources": ["x", "x"]})],
    [field("x"), field("y", "computed", computed={"operation": "multiply", "sources": ["x", "x"]})],
])
def test_dependencies_rejected(fields):
    with pytest.raises(ValidationError):
        generate_custom(request(fields))


def test_template_lifecycle_and_persistence(client, tmp_path):
    template = request()["template"]
    response = client.post(f"{BASE}/templates", json=template)
    assert response.status_code == 201
    saved = response.json()
    id = saved["id"]
    assert client.get(f"{BASE}/templates/{id}").json() == saved
    assert len(client.get(f"{BASE}/templates").json()) == 1
    # A fresh repository/connection reads the same on-disk definition.
    reopened = TemplateStore(tmp_path / "templates.sqlite3")
    assert reopened.get(id).name == template["name"]
    template["name"] = "Updated"
    assert client.put(f"{BASE}/templates/{id}", json=template).json()["name"] == "Updated"
    duplicate = client.post(f"{BASE}/templates/{id}/duplicate").json()
    assert duplicate["id"] != id
    exported = client.get(f"{BASE}/templates/{id}/export")
    imported = client.post(f"{BASE}/templates/import", json=exported.json())
    assert imported.status_code == 201 and imported.json()["id"] != id
    assert imported.json()["fields"] == exported.json()["fields"]
    assert client.delete(f"{BASE}/templates/{id}").status_code == 204
    assert client.get(f"{BASE}/templates/{id}").status_code == 404
    assert client.put(f"{BASE}/templates/{id}", json=template).status_code == 404
    assert client.post(f"{BASE}/templates/import", json={"fields": []}).status_code == 422


@pytest.mark.parametrize("format", ["csv", "json", "xlsx"])
def test_exports_match_generation_and_safe_formulas(client, format):
    payload = request([field("formula", default="=HYPERLINK(\"example\")", options={"max_length": 100}), field("id", "sequence", unique=True)])
    dataset = client.post(f"{BASE}/custom/generate", json=payload).json()
    response = client.post(f"{BASE}/custom/export?format={format}", json=payload)
    assert response.status_code == 200
    if format == "json":
        assert response.json() == dataset
    elif format == "csv":
        rows = list(csv.DictReader(StringIO(response.content.decode("utf-8-sig"))))
        assert len(rows) == 10 and rows[0]["formula"].startswith("'=")
        assert rows[0]["id"] == dataset["data"][0]["id"]
    else:
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.active.max_row == 11
        assert workbook.active["A2"].data_type == "s"
        assert workbook.active["A2"].value.startswith("'=")
        assert workbook.active["B2"].value == dataset["data"][0]["id"]
        workbook.close()


def test_schema_and_body_limits(client):
    assert client.post(f"{BASE}/custom/generate", json=request([field(f"f{i}") for i in range(51)])).status_code == 422
    assert client.post(f"{BASE}/templates/import", content=b" " * 262145, headers={"Content-Type": "application/json"}).status_code == 413
    assert client.post(f"{BASE}/custom/preview", json=request()).status_code == 200
    assert client.post(f"{BASE}/custom/validate", json=request()).json()["valid"] is True


def test_nullable_unique_boolean():
    rows = generate_custom(request([field(type="boolean", unique=True, required=False, null_probability=1)], count=100)).data
    assert all(r["value"] is None for r in rows)


def test_large_permitted_request():
    result = generate_custom(request([field(f"field_{i}", "sequence", unique=True) for i in range(50)], count=1000))
    assert len(result.data) == 1000 and len(result.data[0]) == 50


def test_unique_dates_full_domain():
    data = generate_custom(request([field(type="date", unique=True, options={"min_date": "2026-01-01", "max_date": "2026-01-10"})], count=10)).data
    assert len({r["value"] for r in data}) == 10


@pytest.mark.parametrize("value", [0.29, 0.07, -0.29, 1.23])
def test_decimal_boundaries_do_not_drift(value):
    data = generate_custom(request([field(type="decimal", options={"min": value, "max": value, "precision": 2})])).data
    assert all(r["value"] == value for r in data)


def test_unknown_domain_retries_are_bounded(monkeypatch):
    from app.modules.data_factory.field_registry import FieldGenerator
    monkeypatch.setitem(REGISTRY, "stuck", FieldGenerator(lambda c, f, i: "same"))
    with pytest.raises(ValueError, match="bounded attempts"):
        generate_custom(request([field(type="stuck", unique=True)]))


def test_computed_null_and_unique_errors():
    fields = [field("x", "integer", required=False, null_probability=1), field("total", "computed", required=False, computed={"operation": "multiply", "sources": ["x", "x"]})]
    assert all(r["total"] is None for r in generate_custom(request(fields)).data)
    fields[1]["required"] = True
    with pytest.raises(ValidationError, match="required sources"):
        generate_custom(request(fields))
    fields[0] = field("x", "integer", default=1)
    fields[1]["unique"] = True
    with pytest.raises(ValueError, match="duplicate computed"):
        generate_custom(request(fields))


def test_output_limit():
    fields = [field(f"f{i}", options={"min_length": 256, "max_length": 256}) for i in range(50)]
    with pytest.raises(ValueError, match="output limit"):
        generate_custom(request(fields, count=1000))


def test_store_unavailable(client, monkeypatch):
    monkeypatch.setenv("DATA_FACTORY_DB", __file__)
    response = client.get(f"{BASE}/templates")
    assert response.status_code == 503
    assert response.json()["detail"] == "Template storage is unavailable"
