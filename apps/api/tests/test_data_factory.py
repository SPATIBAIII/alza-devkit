import csv
import json
import random
from io import BytesIO, StringIO

import pytest
from faker import Faker
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.modules.data_factory.exporters import FIELDS, export_households
from app.modules.data_factory.generator import generate_households
from app.modules.data_factory.schemas import HouseholdDataset

client = TestClient(app)
URL = "/api/v1/data-factory/households"


@pytest.mark.parametrize("count", [5, 100, 1000])
def test_generation(count):
    records = generate_households(count)
    assert len(records) == count
    assert len({row["household_id"] for row in records}) == count
    assert records[0]["household_id"] == "HH-00001"
    assert records[-1]["household_id"] == f"HH-{count:05d}"
    assert all(1 <= row["household_size"] <= 12 for row in records)
    for row in records:
        assert "Fictional" in row["barangay"]
        income = row["monthly_income"]
        assert 3000 <= income <= 80000
        assert row["income_category"] == ("Low" if income < 12000 else "Middle" if income < 40000 else "High")


def test_seed_and_global_state():
    state = random.getstate()
    faker_state = Faker().random.getstate()
    assert generate_households(seed=0) == generate_households(seed=0)
    assert generate_households(seed=0) != generate_households(seed=42)
    assert random.getstate() == state
    assert Faker().random.getstate() == faker_state


@pytest.mark.parametrize("count", [0, -1, 1001, 1.5, True])
def test_generator_rejects_invalid_count(count):
    with pytest.raises(ValueError, match="Record count"):
        generate_households(count=count)


@pytest.mark.parametrize("seed", [-1, 1.5, True])
def test_generator_rejects_invalid_seed(seed):
    with pytest.raises(ValueError, match="Seed"):
        generate_households(seed=seed)


@pytest.mark.parametrize("suffix", ["", "/export"])
@pytest.mark.parametrize("query", ["count=0", "count=-1", "count=1001", "count=1.5", "count=abc", "seed=-1", "seed=1.5", "seed=abc"])
def test_invalid_queries(suffix, query):
    response = client.get(f"{URL}{suffix}?{query}")
    assert response.status_code == 422
    assert response.json()["detail"]


def test_defaults_and_health():
    response = client.get(URL)
    assert response.status_code == 200
    body = response.json()
    assert body == {"success": True, "dataset": "households", "total_records": 100, "seed": 42, "synthetic": True, "data": generate_households()}
    assert client.get("/api/v1/health").json()["status"] == "healthy"


@pytest.mark.parametrize("format,mime", [("csv", "text/csv"), ("json", "application/json"), ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")])
def test_exports_match_preview(format, mime):
    preview = client.get(f"{URL}?count=5&seed=7").json()
    response = client.get(f"{URL}/export?count=5&seed=7&format={format}")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(mime)
    assert response.headers["content-disposition"] == f'attachment; filename="synthetic-households-5-seed-7.{format}"'
    if format == "json":
        assert json.loads(response.content) == preview
    elif format == "csv":
        rows = list(csv.DictReader(StringIO(response.content.decode("utf-8-sig"))))
        assert rows == [{key: str(value) for key, value in row.items()} for row in preview["data"]]
    else:
        workbook = load_workbook(BytesIO(response.content))
        rows = list(workbook.active.values)
        assert rows[0] == FIELDS
        assert rows[1:] == [tuple(row[field] for field in FIELDS) for row in preview["data"]]
        workbook.close()


@pytest.mark.parametrize("value", ["=1+1", "+SUM(A1:A2)", "-1+1", "@SUM(A1)", "\t=1", "\r=1", "\n=1", "  =1", "\x01=1", "\ufeff=1"])
def test_formula_injection(value):
    records = generate_households(1)
    records[0]["household_head"] = value
    records[0]["barangay"] = value
    dataset = HouseholdDataset(total_records=1, seed=42, data=records)
    rows = list(csv.DictReader(StringIO(export_households(dataset, "csv").decode("utf-8-sig"), newline="")))
    assert rows[0]["household_head"] == "'" + value
    assert rows[0]["barangay"] == "'" + value
    workbook = load_workbook(BytesIO(export_households(dataset, "xlsx")))
    assert workbook.active["B2"].data_type == "s"
    expected_excel = "'" + value.replace("\x01", "").replace("\r", "\n")
    assert workbook.active["B2"].value == expected_excel
    assert workbook.active["C2"].data_type == "s"
    assert workbook.active["C2"].value == expected_excel
    workbook.close()
    assert json.loads(export_households(dataset, "json"))["data"][0]["household_head"] == value


def test_unicode_and_csv_quoting():
    records = generate_households(1)
    records[0]["household_head"] = 'Peña, "Test"\nSynthetic'
    dataset = HouseholdDataset(total_records=1, seed=42, data=records)
    rows = list(csv.DictReader(StringIO(export_households(dataset, "csv").decode("utf-8-sig"), newline="")))
    assert rows[0]["household_head"] == records[0]["household_head"]


def test_export_format_validation():
    assert client.get(f"{URL}/export?format=pdf").status_code == 422


def test_cors():
    response = client.get(URL, headers={"Origin": "http://localhost:3000"})
    assert response.headers["access-control-allow-origin"] == "http://localhost:3000"
    response = client.get(URL, headers={"Origin": "https://untrusted.example"})
    assert "access-control-allow-origin" not in response.headers
